"""
知识库自学习系统 - 从用户文档中自动学习

这是一个库模块，供 Agent 通过 python_run 调用。
不注册为工具，符合 Claude 极简工具设计理念。

使用方式（python_run）：
```python
from backend.tools.base.knowledge_learning import get_knowledge_learner
learner = get_knowledge_learner()
stats = learner.scan_and_learn()
print(stats)
```

学习产出：
1. DocMap - 文档结构映射
2. Ontology - 领域本体
3. Skills - 自动生成的技能（SKILL.md）
4. KnowledgeGraph - 知识图谱实体和关系

存储位置：knowledge_base/learned/
"""

import os
import re
import json
import time
import hashlib
import logging
import threading
from pathlib import Path
from collections import Counter
from typing import Dict, List, Any, Optional, Tuple
from dataclasses import dataclass, field, asdict
from datetime import datetime

logger = logging.getLogger(__name__)

# 本地模块
try:
    from .knowledge_graph import (
        get_knowledge_graph, 
        get_extractor,
        EntityType, 
        RelationType,
    )
    _HAS_KG = True
except ImportError:
    _HAS_KG = False

try:
    from .learning_middleware import learn_from_document, get_learning_manager
    _HAS_LEARNING = True
except ImportError:
    _HAS_LEARNING = False

try:
    from backend.engine.agent.model_manager import get_model_manager
    _HAS_MODEL_MANAGER = True
except ImportError:
    _HAS_MODEL_MANAGER = False


# ============================================================
# 配置
# ============================================================
PROJECT_ROOT = Path(__file__).parent.parent.parent.parent
KNOWLEDGE_BASE = PROJECT_ROOT / "knowledge_base"
LEARNED_DIR = KNOWLEDGE_BASE / "learned"  # 学习产出目录
USERS_DIR = KNOWLEDGE_BASE / "users"
LEARNED_TEMPLATES_DIR = LEARNED_DIR / "templates"

@dataclass
class DocumentMeta:
    """文档元数据"""
    path: str
    hash: str  # 内容哈希，用于检测变更
    size: int
    processed_at: str
    doc_type: str  # pdf, docx, xlsx, txt, md
    domain: str  # 推断的领域
    structure: Dict = field(default_factory=dict)  # DocMap


@dataclass
class LearnedSkill:
    """学习到的技能"""
    name: str
    description: str
    triggers: List[str]  # 触发词
    workflow: List[str]  # 工作流步骤
    tools: List[str]  # 使用的工具
    source_docs: List[str]  # 来源文档
    confidence: float = 0.5
    created_at: str = field(default_factory=lambda: datetime.now().isoformat())
    
    def to_skill_md(self) -> str:
        """转换为 SKILL.md 格式（含 level/domain/source 便于 registry 与 profile 过滤）"""
        domain = self.name[:-9] if self.name.endswith("_analysis") else "learned"
        return f"""---
name: {self.name}
description: {self.description}
level: learned
domain: {domain}
source: learned
triggers: {json.dumps(self.triggers, ensure_ascii=False)}
tools: {json.dumps(self.tools)}
confidence: {self.confidence}
auto_generated: true
---

# {self.name}

{self.description}

## 触发条件

当用户请求涉及以下关键词时使用此技能：
{chr(10).join(f'- {t}' for t in self.triggers)}

## 工作流程

{chr(10).join(f'{i+1}. {step}' for i, step in enumerate(self.workflow))}

## 来源文档

{chr(10).join(f'- {doc}' for doc in self.source_docs)}
"""


class KnowledgeLearner:
    """
    知识库自学习器
    
    功能：
    1. 扫描用户文档目录
    2. 解析文档内容
    3. 提取结构化知识
    4. 生成 DocMap、Ontology、Skills
    5. 更新知识图谱
    """
    
    def __init__(self, watch_dirs: Optional[List[str]] = None):
        """
        Args:
            watch_dirs: 要监控的目录列表
        """
        LEARNED_DIR.mkdir(parents=True, exist_ok=True)
        self.watch_dirs = watch_dirs or [
            str(KNOWLEDGE_BASE / "domain"),
            str(KNOWLEDGE_BASE / "global"),
            str(PROJECT_ROOT / "tmp" / "inputs"),  # 用户上传目录
            str(LEARNED_DIR / "web_cache"),  # web_search 结果缓存，供检索复用
        ]
        
        # 已处理文档索引
        self._processed_index_path = LEARNED_DIR / "processed_index.json"
        self._processed: Dict[str, DocumentMeta] = {}
        self._load_processed_index()
        
        # 学习到的技能
        self._skills_path = LEARNED_DIR / "learned_skills.json"
        self._skills: Dict[str, LearnedSkill] = {}
        self._load_skills()
        
        # 领域本体
        self._ontology_path = LEARNED_DIR / "domain_ontology.json"
        self._ontology: Dict[str, Any] = {
            "entity_types": {},
            "relation_types": {},
            "domain_terms": {},
        }
        self._load_ontology()

        # 自动学习轮询（轻量 fallback；若未来引入 watchdog 可替换）
        self._auto_watch_thread: Optional[threading.Thread] = None
        self._auto_watch_stop = threading.Event()
        self._auto_watch_interval_sec = 30
    
    def _load_processed_index(self):
        """加载已处理文档索引"""
        if self._processed_index_path.exists():
            try:
                data = json.loads(self._processed_index_path.read_text(encoding="utf-8"))
                for path, meta in data.items():
                    self._processed[path] = DocumentMeta(**meta)
            except Exception as e:
                logger.warning("加载处理索引失败: %s", e)
    
    def _save_processed_index(self):
        """保存已处理文档索引"""
        data = {path: asdict(meta) for path, meta in self._processed.items()}
        self._processed_index_path.write_text(
            json.dumps(data, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
    
    def _load_skills(self):
        """加载学习到的技能"""
        if self._skills_path.exists():
            try:
                data = json.loads(self._skills_path.read_text(encoding="utf-8"))
                for name, skill_data in data.items():
                    self._skills[name] = LearnedSkill(**skill_data)
            except Exception as e:
                logger.warning("加载技能失败: %s", e)
    
    def _save_skills(self):
        """保存学习到的技能"""
        data = {name: asdict(skill) for name, skill in self._skills.items()}
        self._skills_path.write_text(
            json.dumps(data, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
    
    def _load_ontology(self):
        """加载领域本体"""
        if self._ontology_path.exists():
            try:
                self._ontology = json.loads(self._ontology_path.read_text(encoding="utf-8"))
            except Exception as e:
                logger.warning("加载本体失败: %s", e)
    
    def _save_ontology(self):
        """保存领域本体"""
        self._ontology_path.write_text(
            json.dumps(self._ontology, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
    
    def _get_file_hash(self, path: str) -> str:
        """计算文件哈希"""
        h = hashlib.md5()
        with open(path, 'rb') as f:
            for chunk in iter(lambda: f.read(8192), b''):
                h.update(chunk)
        return h.hexdigest()
    
    def _infer_doc_type(self, path: str) -> str:
        """推断文档类型"""
        ext = Path(path).suffix.lower()
        type_map = {
            '.pdf': 'pdf',
            '.docx': 'docx',
            '.doc': 'doc',
            '.xlsx': 'xlsx',
            '.xls': 'xls',
            '.txt': 'txt',
            '.md': 'md',
            '.json': 'json',
            '.csv': 'csv',
        }
        return type_map.get(ext, 'unknown')
    
    def _infer_domain(self, path: str, content: str) -> str:
        """推断文档领域"""
        path_lower = path.lower()
        content_lower = content[:2000].lower()
        
        # 基于路径和内容的简单规则
        domain_keywords = {
            "bidding": ["招标", "投标", "标书", "评标", "废标", "bid"],
            "contract": ["合同", "协议", "条款", "contract", "agreement"],
            "finance": ["财务", "会计", "报表", "finance", "accounting"],
            "legal": ["法律", "法规", "条例", "legal", "law"],
            "technical": ["技术", "规范", "标准", "technical", "spec"],
            "management": ["管理", "流程", "制度", "management", "process"],
        }
        
        scores = {}
        for domain, keywords in domain_keywords.items():
            score = sum(1 for kw in keywords if kw in path_lower or kw in content_lower)
            scores[domain] = score
        
        if max(scores.values()) > 0:
            return max(scores, key=scores.get)
        return "general"
    
    def _extract_docmap(self, content: str, doc_type: str) -> Dict:
        """提取文档结构映射（DocMap）"""
        docmap = {
            "sections": [],
            "keywords": [],
            "summary": "",
        }
        
        lines = content.split('\n')

        # 提取章节（基于标题模式）；每组 pattern 对应 (group_index_for_title, level 或 level 推断方式)
        section_patterns = [
            (r'^(#{1,6})\s+(.+)$', 2, None),  # Markdown: group(1)=hashes, group(2)=title, level=len(hashes)
            (r'^第[一二三四五六七八九十\d]+[章节条款]\s*(.+)$', 1, 1),  # 中文章节: group(1)=title, level=1
            (r'^(\d+\.)+\s*(.+)$', 2, None),  # 编号标题: group(2)=title，level 由点数推断
        ]
        for i, line in enumerate(lines):
            line_stripped = line.strip()
            for pattern, title_group, fixed_level in section_patterns:
                match = re.match(pattern, line_stripped)
                if match:
                    title = match.group(title_group).strip() if match.lastindex >= title_group else line_stripped
                    if fixed_level is not None:
                        level = fixed_level
                    elif pattern.startswith(r'^(#{1,6})'):
                        level = len(match.group(1))
                    else:
                        # 编号标题：1. -> 1, 1.2. -> 2, 1.2.3. -> 3
                        num_part = match.group(0).split()[0] if match.groups() else ''
                        level = min(6, max(1, num_part.rstrip('.').count('.') + 1))
                    docmap["sections"].append({"title": title, "line": i + 1, "level": level})
                    break

        # 提取关键词：2–6 字中文片段，过滤常见停用词
        _stop = frozenset({
            "的是", "可以", "进行", "通过", "以及", "如果", "因为", "所以", "这个", "那个",
            "一种", "没有", "不是", "什么", "如何", "我们", "他们", "它们", "已经", "或者",
            "但是", "然而", "因此", "其中", "这些", "那些", "这样", "那样", "为了", "由于",
            "并且", "而且", "虽然", "尽管", "应当", "即使", "无论", "不管", "关于", "对于",
        })
        words = re.findall(r'[\u4e00-\u9fff]{2,6}', content)
        word_counts = Counter(words)
        docmap["keywords"] = [w for w, c in word_counts.most_common(30) if c >= 2 and w not in _stop][:20]
        
        # 生成摘要（前 200 字）
        docmap["summary"] = content[:200].replace('\n', ' ').strip()
        
        return docmap
    
    def _extract_workflow(self, content: str) -> List[str]:
        """从文档中提取工作流程"""
        workflows = []

        # 模式1: 编号步骤
        steps = re.findall(r'(?:步骤|Step)\s*[\d一二三四五六七八九十]+[.:：]?\s*(.+)', content)
        if steps:
            workflows.extend(steps[:10])
        
        # 模式2: 流程关键词
        process_keywords = ["首先", "然后", "接着", "最后", "第一", "第二"]
        for kw in process_keywords:
            matches = re.findall(f'{kw}[，,]?(.{{10,50}})[。.]', content)
            workflows.extend(matches[:3])
        
        return list(set(workflows))[:10]

    def _extract_style_profile(self, content: str, doc_type: str) -> Dict[str, Any]:
        """提取用户文档风格画像（用于格式/语气偏好约束）。"""
        lines = [ln.strip() for ln in content.splitlines() if ln.strip()]
        title_lines = [ln for ln in lines if len(ln) <= 24 and (ln.startswith("#") or re.match(r"^[一二三四五六七八九十\d]+[、.．]", ln))]
        bullet_lines = [ln for ln in lines if re.match(r"^[-*•]\s+", ln)]
        avg_sentence_len = 0.0
        sentences = re.split(r"[。！？!?\n]+", content)
        valid_sentences = [s.strip() for s in sentences if s.strip()]
        if valid_sentences:
            avg_sentence_len = round(sum(len(s) for s in valid_sentences) / len(valid_sentences), 2)
        top_terms = [w for w, c in Counter(re.findall(r"[\u4e00-\u9fff]{2,8}", content)).most_common(15) if c >= 2]
        return {
            "doc_type": doc_type,
            "style_signals": {
                "has_headings": bool(title_lines),
                "has_bullets": bool(bullet_lines),
                "avg_sentence_len": avg_sentence_len,
                "paragraph_count": max(1, len(lines)),
            },
            "preferred_terms": top_terms[:12],
            "sample_headings": title_lines[:8],
        }

    def _extract_document_template(self, content: str, doc_type: str) -> Dict[str, Any]:
        """提取可复用模板骨架（章节和占位字段）。"""
        sections = []
        for line in content.splitlines():
            s = line.strip()
            if not s:
                continue
            if s.startswith("#") or re.match(r"^[一二三四五六七八九十\d]+[、.．]", s):
                sections.append(s)
        placeholders = sorted(set(re.findall(r"[\[【](.+?)[\]】]", content)))[:30]
        return {
            "doc_type": doc_type,
            "sections": sections[:30],
            "placeholders": placeholders,
            "template_excerpt": content[:1200],
        }

    def _persist_user_style_and_template(
        self,
        path: str,
        domain: str,
        style_profile: Dict[str, Any],
        template_profile: Dict[str, Any],
    ) -> Dict[str, str]:
        """保存用户风格画像与模板提取产物。"""
        user_id = os.getenv("MAIBOT_USER_ID", "default")
        user_dir = USERS_DIR / user_id
        user_dir.mkdir(parents=True, exist_ok=True)
        LEARNED_TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)

        style_path = user_dir / "style_profile.json"
        existing = {}
        if style_path.exists():
            try:
                existing = json.loads(style_path.read_text(encoding="utf-8"))
            except Exception:
                existing = {}
        records = existing.get("records", []) if isinstance(existing, dict) else []
        records.append(
            {
                "source_path": path,
                "domain": domain,
                "updated_at": datetime.now().isoformat(),
                "style_profile": style_profile,
            }
        )
        merged = {
            "user_id": user_id,
            "updated_at": datetime.now().isoformat(),
            "records": records[-80:],
        }
        style_path.write_text(json.dumps(merged, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

        template_name = f"{Path(path).stem}_template.json"
        template_path = LEARNED_TEMPLATES_DIR / template_name
        template_payload = {
            "source_path": path,
            "domain": domain,
            "created_at": datetime.now().isoformat(),
            "template": template_profile,
        }
        template_path.write_text(json.dumps(template_payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        return {"style_profile_path": str(style_path), "template_path": str(template_path)}

    def _get_llm_callable(self):
        """获取用于实体关系抽取的 LLM 调用函数。
        
        仅在 KNOWLEDGE_LLM_EXTRACT=1 环境变量显式开启时才调用 LLM，
        默认关闭以避免后台任务占用本地模型资源影响主聊天响应。
        """
        if not _HAS_MODEL_MANAGER:
            return None
        # 默认关闭：LLM 辅助抽取需显式开启（KNOWLEDGE_LLM_EXTRACT=1）
        if not os.environ.get("KNOWLEDGE_LLM_EXTRACT", "").lower() in ("1", "true", "yes"):
            return None
        try:
            llm = get_model_manager().create_llm(task_type="analysis")
        except Exception as e:
            logger.warning("初始化本体提取 LLM 失败: %s", e)
            return None

        def _call(prompt: str) -> str:
            try:
                resp = llm.invoke(prompt)
                content = getattr(resp, "content", "")
                if isinstance(content, str):
                    return content
                if isinstance(content, list):
                    chunks: List[str] = []
                    for item in content:
                        if isinstance(item, dict):
                            chunks.append(str(item.get("text", "")))
                        else:
                            chunks.append(str(item))
                    return "\n".join([c for c in chunks if c])
                return str(content)
            except Exception as invoke_err:
                logger.warning("LLM 调用失败，回退规则抽取: %s", invoke_err)
                return ""

        return _call
    
    def scan_and_learn(self, force: bool = False) -> Dict[str, Any]:
        """
        扫描目录并学习
        
        Args:
            force: 是否强制重新处理所有文档
        
        Returns:
            学习结果统计
        """
        stats = {
            "scanned": 0,
            "new": 0,
            "updated": 0,
            "skipped": 0,
            "errors": 0,
            "skills_generated": 0,
            "entities_extracted": 0,
        }
        
        # 扫描所有目录
        for watch_dir in self.watch_dirs:
            if not os.path.exists(watch_dir):
                continue
            
            for root, dirs, files in os.walk(watch_dir):
                # 跳过隐藏目录
                dirs[:] = [d for d in dirs if not d.startswith('.')]
                
                for file in files:
                    if file.startswith('.'):
                        continue
                    
                    path = os.path.join(root, file)
                    if "web_cache" in path:
                        if os.path.getmtime(path) < time.time() - 7 * 24 * 3600:
                            continue
                    doc_type = self._infer_doc_type(path)
                    
                    if doc_type == 'unknown':
                        continue
                    
                    stats["scanned"] += 1
                    
                    try:
                        result = self.learn_from_document(path, force=force)
                        if result["status"] == "new":
                            stats["new"] += 1
                        elif result["status"] == "updated":
                            stats["updated"] += 1
                        else:
                            stats["skipped"] += 1
                        
                        stats["entities_extracted"] += result.get("entities", 0)
                        
                    except Exception as e:
                        stats["errors"] += 1
                        logger.warning("处理 %s 失败: %s", path, e)
        
        # 生成技能
        new_skills = self._generate_skills_from_knowledge()
        stats["skills_generated"] = len(new_skills)
        
        # 保存所有数据
        self._save_processed_index()
        self._save_skills()
        self._save_ontology()
        
        return stats
    
    def learn_from_document(self, path: str, force: bool = False) -> Dict[str, Any]:
        """
        从单个文档学习
        
        Args:
            path: 文档路径
            force: 是否强制重新处理
        
        Returns:
            学习结果
        """
        # 检查是否已处理
        file_hash = self._get_file_hash(path)
        
        if not force and path in self._processed:
            if self._processed[path].hash == file_hash:
                return {"status": "skipped", "reason": "unchanged"}
        
        # 读取文档内容
        content = self._read_document(path)
        if not content:
            return {"status": "error", "reason": "cannot_read"}
        
        # 推断类型和领域
        doc_type = self._infer_doc_type(path)
        domain = self._infer_domain(path, content)
        style_profile = self._extract_style_profile(content, doc_type)
        template_profile = self._extract_document_template(content, doc_type)
        persisted_paths = self._persist_user_style_and_template(
            path=path,
            domain=domain,
            style_profile=style_profile,
            template_profile=template_profile,
        )
        
        # 提取 DocMap
        docmap = self._extract_docmap(content, doc_type)
        
        # 提取实体和关系（使用知识图谱）
        entities_count = 0
        if _HAS_LEARNING:
            try:
                result = learn_from_document(
                    task_id=f"doc_{file_hash[:8]}",
                    document_text=content[:10000],  # 限制长度
                    document_source=path,
                )
                entities_count = result.get("entities_count", 0)
            except Exception as e:
                logger.warning("知识提取失败: %s", e)
        
        # 更新本体（含 learned/ontology 实体与关系）
        self._update_ontology(content, domain, source=path)
        
        # 保存元数据
        meta = DocumentMeta(
            path=path,
            hash=file_hash,
            size=os.path.getsize(path),
            processed_at=datetime.now().isoformat(),
            doc_type=doc_type,
            domain=domain,
            structure=docmap,
        )
        
        status = "updated" if path in self._processed else "new"
        self._processed[path] = meta
        
        return {
            "status": status,
            "domain": domain,
            "sections": len(docmap.get("sections", [])),
            "keywords": len(docmap.get("keywords", [])),
            "entities": entities_count,
            "style_profile_path": persisted_paths.get("style_profile_path", ""),
            "template_path": persisted_paths.get("template_path", ""),
        }
    
    def _read_document(self, path: str) -> Optional[str]:
        """读取文档内容"""
        doc_type = self._infer_doc_type(path)
        
        try:
            if doc_type in ['txt', 'md', 'json', 'csv']:
                return Path(path).read_text(encoding='utf-8')
            
            elif doc_type == 'pdf':
                try:
                    import pdfplumber
                    with pdfplumber.open(path) as pdf:
                        return '\n'.join(page.extract_text() or '' for page in pdf.pages)
                except ImportError:
                    return None
            
            elif doc_type == 'docx':
                try:
                    import docx
                    doc = docx.Document(path)
                    return '\n'.join(para.text for para in doc.paragraphs)
                except ImportError:
                    return None
            
            elif doc_type in ['xlsx', 'xls']:
                try:
                    import pandas as pd
                    df = pd.read_excel(path)
                    return df.to_string()
                except ImportError:
                    return None
            
        except Exception as e:
            logger.warning("读取 %s 失败: %s", path, e)
        
        return None
    
    def _update_ontology(self, content: str, domain: str, source: str = ""):
        """更新领域本体；同时将实体与关系写入 learned/ontology/（通过 KnowledgeGraph）。"""
        # 提取专业术语（简单规则）
        # 中文专业术语通常是 2-6 字的名词短语
        terms = re.findall(r'[\u4e00-\u9fff]{2,6}', content)
        
        if domain not in self._ontology["domain_terms"]:
            self._ontology["domain_terms"][domain] = {}
        
        for term in terms:
            if term not in self._ontology["domain_terms"][domain]:
                self._ontology["domain_terms"][domain][term] = 0
            self._ontology["domain_terms"][domain][term] += 1
        
        # 同时写入 learned/ontology/（entities/relations）供图谱与 API 使用
        if _HAS_KG:
            try:
                extractor = get_extractor()
                # 先规则抽取，保证稳定性
                entities = extractor.extract_entities_simple(content, source=source)
                relations = extractor.extract_relations_simple(content, entities, source=source)

                # 再尝试 LLM 抽取补充（Schema-Driven），提升召回与准确率
                llm_call = self._get_llm_callable()
                if llm_call is not None:
                    llm_entities, llm_relations = extractor.extract_with_llm(
                        content,
                        llm_func=llm_call,
                        source=source,
                        domain=domain,
                    )
                    entities.extend(llm_entities or [])
                    relations.extend(llm_relations or [])

                extractor.accumulate(entities, relations)
            except Exception as e:
                logger.warning("本体写入 learned/ontology 失败: %s", e)
    
    def _generate_skills_from_knowledge(self) -> List[LearnedSkill]:
        """从积累的知识中生成技能"""
        new_skills = []
        
        # 按领域分组文档
        domains = {}
        for path, meta in self._processed.items():
            domain = meta.domain
            if domain not in domains:
                domains[domain] = []
            domains[domain].append(meta)
        
        # 为每个领域生成技能
        for domain, metas in domains.items():
            if len(metas) < 2:  # 至少 2 个文档才生成技能
                continue
            
            # 合并关键词
            all_keywords = []
            all_sections = []
            source_docs = []
            
            for meta in metas:
                all_keywords.extend(meta.structure.get("keywords", []))
                all_sections.extend([s["title"] for s in meta.structure.get("sections", [])])
                source_docs.append(meta.path)
            
            # 取高频关键词作为触发词
            keyword_counts = Counter(all_keywords)
            triggers = [kw for kw, c in keyword_counts.most_common(10) if c >= 2]
            
            if not triggers:
                continue
            
            # 生成技能
            skill_name = f"{domain}_analysis"
            
            if skill_name not in self._skills:
                skill = LearnedSkill(
                    name=skill_name,
                    description=f"自动学习的{domain}领域分析技能，基于 {len(metas)} 个文档",
                    triggers=triggers[:10],
                    workflow=[
                        f"1. 识别{domain}领域的关键要素",
                        "2. 使用 search_knowledge 检索相关知识",
                        "3. 使用 python_run 进行数据分析",
                        "4. 生成分析报告",
                    ],
                    tools=["search_knowledge", "python_run", "read_file"],
                    source_docs=source_docs[:5],
                    confidence=min(0.9, 0.3 + len(metas) * 0.1),
                )
                
                self._skills[skill_name] = skill
                new_skills.append(skill)
                
                # 生成 SKILL.md 文件
                self._export_skill_md(skill)
        
        return new_skills
    
    def _export_skill_md(self, skill: LearnedSkill):
        """导出技能为 SKILL.md 文件"""
        skill_dir = LEARNED_DIR / "skills" / skill.name
        skill_dir.mkdir(parents=True, exist_ok=True)
        
        skill_md_path = skill_dir / "SKILL.md"
        skill_md_path.write_text(skill.to_skill_md(), encoding='utf-8')
        
        logger.info("生成技能: %s", skill_md_path)
    
    def get_stats(self) -> Dict[str, Any]:
        """获取学习统计"""
        return {
            "processed_documents": len(self._processed),
            "learned_skills": len(self._skills),
            "domains": list(self._ontology.get("domain_terms", {}).keys()),
            "watch_dirs": self.watch_dirs,
        }

    def start_auto_watch(self, interval_sec: int = 30, force: bool = False) -> Dict[str, Any]:
        """启动后台自动学习轮询。"""
        if self._auto_watch_thread and self._auto_watch_thread.is_alive():
            return {"status": "running", "interval_sec": self._auto_watch_interval_sec}

        self._auto_watch_interval_sec = max(5, int(interval_sec))
        self._auto_watch_stop.clear()

        def _loop():
            while not self._auto_watch_stop.is_set():
                try:
                    self.scan_and_learn(force=force)
                except Exception as e:
                    logger.warning("自动学习轮询失败: %s", e)
                self._auto_watch_stop.wait(self._auto_watch_interval_sec)

        self._auto_watch_thread = threading.Thread(
            target=_loop,
            name="knowledge-learner-auto-watch",
            daemon=True,
        )
        self._auto_watch_thread.start()
        return {"status": "started", "interval_sec": self._auto_watch_interval_sec}

    def stop_auto_watch(self) -> Dict[str, Any]:
        """停止后台自动学习轮询。"""
        if not self._auto_watch_thread or not self._auto_watch_thread.is_alive():
            return {"status": "stopped"}
        self._auto_watch_stop.set()
        self._auto_watch_thread.join(timeout=2.0)
        return {"status": "stopped"}


# ============================================================
# 全局实例
# ============================================================
_knowledge_learner: Optional[KnowledgeLearner] = None


def get_knowledge_learner() -> KnowledgeLearner:
    """获取知识学习器单例"""
    global _knowledge_learner
    if _knowledge_learner is None:
        _knowledge_learner = KnowledgeLearner()
    return _knowledge_learner


# ============================================================
# 便捷函数（供 python_run 调用）
# ============================================================
def scan_and_learn(directory: Optional[str] = None, force: bool = False) -> Dict:
    """
    扫描目录并学习文档知识
    
    Args:
        directory: 要扫描的目录（None 使用默认目录）
        force: 是否强制重新处理
    
    Returns:
        学习统计
    
    Example:
        from backend.tools.base.knowledge_learning import scan_and_learn
        stats = scan_and_learn("/path/to/docs")
    """
    learner = get_knowledge_learner()
    if directory:
        saved = learner.watch_dirs
        try:
            learner.watch_dirs = [directory]
            return learner.scan_and_learn(force=force)
        finally:
            learner.watch_dirs = saved
    return learner.scan_and_learn(force=force)


def learn_document(path: str) -> Dict:
    """
    学习单个文档
    
    Args:
        path: 文档路径
    
    Returns:
        学习结果
    
    Example:
        from backend.tools.base.knowledge_learning import learn_document
        result = learn_document("/path/to/doc.pdf")
    """
    learner = get_knowledge_learner()
    return learner.learn_from_document(path, force=True)


def get_learning_stats() -> Dict:
    """
    获取学习系统统计
    
    Example:
        from backend.tools.base.knowledge_learning import get_learning_stats
        stats = get_learning_stats()
    """
    return get_knowledge_learner().get_stats()


def start_auto_learning(interval_sec: int = 30, force: bool = False) -> Dict:
    """
    启动自动学习轮询（后台线程）。
    """
    return get_knowledge_learner().start_auto_watch(interval_sec=interval_sec, force=force)


def stop_auto_learning() -> Dict:
    """
    停止自动学习轮询。
    """
    return get_knowledge_learner().stop_auto_watch()


__all__ = [
    # 核心类
    "KnowledgeLearner",
    "LearnedSkill",
    "DocumentMeta",
    # 单例获取
    "get_knowledge_learner",
    # 便捷函数（供 python_run 调用）
    "scan_and_learn",
    "learn_document",
    "get_learning_stats",
    "start_auto_learning",
    "stop_auto_learning",
]
